# Copyright (c) 2025, antruong and contributors
# For license information, please see license.txt

"""Document parser service for extracting text from uploaded files.

Supports .docx, .xlsx, .pdf, .txt formats via Frappe's file system.
"""

import os

import frappe
from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader


# Allowed file extensions for upload validation
ALLOWED_EXTENSIONS = {".docx", ".xlsx", ".pdf", ".txt"}
# Maximum file size: 10MB
MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024


def validate_file_extension(file_path):
	"""Validate that the file has an allowed extension.

	Args:
		file_path: Path to the file to validate.

	Raises:
		frappe.ValidationError: If file extension is not in allowed list.
	"""
	ext = os.path.splitext(file_path)[1].lower()
	if ext not in ALLOWED_EXTENSIONS:
		frappe.throw(
			f"Loại file không được hỗ trợ: {ext}. "
			f"Chỉ chấp nhận: {', '.join(ALLOWED_EXTENSIONS)}"
		)


def get_file_path_from_url(file_url):
	"""Convert a Frappe file URL to an absolute file path.

	Args:
		file_url: Frappe file URL (e.g., /files/myfile.docx).

	Returns:
		Absolute path to the file on disk.

	Raises:
		frappe.ValidationError: If file does not exist.
	"""
	if not file_url:
		frappe.throw("URL file không được để trống")

	# Handle both /files/ and /private/files/ paths
	site_path = frappe.get_site_path()

	if file_url.startswith("/private/files/"):
		file_path = os.path.join(site_path, file_url.lstrip("/"))
	elif file_url.startswith("/files/"):
		file_path = os.path.join(site_path, "public", file_url.lstrip("/"))
	else:
		frappe.throw(f"Đường dẫn file không hợp lệ: {file_url}")

	# Resolve the path and ensure it stays within the site directory
	resolved_path = os.path.realpath(file_path)
	resolved_site = os.path.realpath(site_path)
	if not resolved_path.startswith(resolved_site + os.sep):
		frappe.throw("Truy cập file không hợp lệ")

	if not os.path.exists(resolved_path):
		frappe.throw(f"File không tồn tại: {file_url}")

	# Validate file size
	file_size = os.path.getsize(resolved_path)
	if file_size > MAX_FILE_SIZE_BYTES:
		frappe.throw(
			f"File quá lớn ({file_size / 1024 / 1024:.1f}MB). "
			f"Giới hạn: {MAX_FILE_SIZE_BYTES / 1024 / 1024:.0f}MB"
		)

	return resolved_path


def parse_docx(file_path):
	"""Extract text content from a .docx file.

	Extracts text from paragraphs and tables.

	Args:
		file_path: Absolute path to the .docx file.

	Returns:
		String containing extracted text content.
	"""
	doc = DocxDocument(file_path)
	content_parts = []

	# Extract paragraphs
	for para in doc.paragraphs:
		text = para.text.strip()
		if text:
			content_parts.append(text)

	# Extract tables
	for table in doc.tables:
		table_rows = []
		for row in table.rows:
			cells = [cell.text.strip() for cell in row.cells]
			# Deduplicate adjacent cells with same content (merged cells)
			deduped = []
			prev = None
			for c in cells:
				if c != prev:
					deduped.append(c)
				prev = c
			table_rows.append(" | ".join(deduped))
		if table_rows:
			content_parts.append("\n".join(table_rows))

	return "\n\n".join(content_parts)


def parse_xlsx(file_path):
	"""Extract text content from an .xlsx file.

	Reads all sheets and extracts cell values.

	Args:
		file_path: Absolute path to the .xlsx file.

	Returns:
		String containing extracted text content.
	"""
	wb = load_workbook(file_path, read_only=True, data_only=True)
	content_parts = []

	for sheet_name in wb.sheetnames:
		ws = wb[sheet_name]
		sheet_content = [f"=== Sheet: {sheet_name} ==="]

		for row in ws.iter_rows(values_only=True):
			cells = []
			for cell in row:
				if cell is not None:
					cells.append(str(cell).strip())
			if cells:
				sheet_content.append(" | ".join(cells))

		if len(sheet_content) > 1:
			content_parts.append("\n".join(sheet_content))

	wb.close()
	return "\n\n".join(content_parts)


def parse_pdf(file_path):
	"""Extract text content from a .pdf file.

	Args:
		file_path: Absolute path to the .pdf file.

	Returns:
		String containing extracted text content.
	"""
	reader = PdfReader(file_path)
	content_parts = []

	for page in reader.pages:
		text = page.extract_text()
		if text and text.strip():
			content_parts.append(text.strip())

	return "\n\n".join(content_parts)


def parse_txt(file_path):
	"""Read text content from a .txt file.

	Args:
		file_path: Absolute path to the .txt file.

	Returns:
		String containing the file content.
	"""
	with open(file_path, encoding="utf-8") as f:
		return f.read()


def parse_file(file_url):
	"""Parse an uploaded file and extract text content.

	Main entry point for document parsing. Routes to the appropriate
	parser based on file extension.

	Args:
		file_url: Frappe file URL to parse.

	Returns:
		String containing the extracted text content.

	Raises:
		frappe.ValidationError: If file type is unsupported.
	"""
	file_path = get_file_path_from_url(file_url)
	validate_file_extension(file_path)

	ext = os.path.splitext(file_path)[1].lower()

	parsers = {
		".docx": parse_docx,
		".xlsx": parse_xlsx,
		".pdf": parse_pdf,
		".txt": parse_txt,
	}

	parser = parsers.get(ext)
	if not parser:
		frappe.throw(f"Không tìm thấy parser cho loại file: {ext}")

	content = parser(file_path)

	if not content or not content.strip():
		frappe.throw(
			f"Không thể trích xuất nội dung từ file. "
			f"Vui lòng kiểm tra file có dữ liệu."
		)

	return content


def parse_scan_pdf(file_url):
	"""Parse a scanned PDF file using OCR API.

	Converts a scanned PDF to Markdown text via external OCR service.

	Args:
		file_url: Frappe file URL to the scanned PDF.

	Returns:
		String containing the Markdown text extracted from the PDF.

	Raises:
		frappe.ValidationError: If OCR processing fails.
	"""
	from cnb_2as.services.ocr_service import ocr_pdf_to_markdown

	file_path = get_file_path_from_url(file_url)

	ext = os.path.splitext(file_path)[1].lower()
	if ext != ".pdf":
		frappe.throw(
			f"File scan phải là PDF. Loại file hiện tại: {ext}"
		)

	return ocr_pdf_to_markdown(file_path)


def extract_fields_from_markdown(markdown_text, doc_type="eval"):
	"""Extract key fields and check completeness from OCR JSON content.

	Parses JSON text (from OCR) to identify employee information fields.
	(Kept the name extract_fields_from_markdown for backward compatibility)

	Args:
		markdown_text: JSON string from OCR.
		doc_type: "eval" for evaluation form, "report" for work report.

	Returns:
		Dict with fields and warnings.
	"""
	import json
	fields = {}
	warnings = []

	if not markdown_text or not markdown_text.strip():
		return {"fields": fields, "warnings": ["Nội dung file rỗng"], "is_complete": False}

	try:
		data = json.loads(markdown_text)
	except json.JSONDecodeError:
		# Fallback if not JSON
		return {"fields": {}, "warnings": ["Lỗi parse JSON từ OCR"], "is_complete": False}

	if doc_type == "eval":
		info = data.get("thong_tin_nhan_vien", {})
		if info.get("ho_ten"): fields["Họ và tên"] = info["ho_ten"]
		if info.get("msnv"): fields["MSNV"] = info["msnv"]
		if info.get("vi_tri"): fields["Vị trí công việc"] = info["vi_tri"]
		if info.get("phong_ban"): fields["Phòng ban"] = info["phong_ban"]
		if info.get("ngay_bat_dau_hd"): fields["Ngày bắt đầu HĐ"] = info["ngay_bat_dau_hd"]
		if info.get("ngay_het_han_hd"): fields["Ngày hết hạn HĐ"] = info["ngay_het_han_hd"]
		if info.get("ngay_nhan_viec"): fields["Ngày nhận việc"] = info["ngay_nhan_viec"]
		
		# Check warnings for required fields
		required = [
			("ho_ten", "Họ và tên CBNV"),
			("msnv", "MSNV"),
			("vi_tri", "Vị trí công việc"),
			("phong_ban", "Phòng ban")
		]
		for key, label in required:
			if not info.get(key) or info.get(key).strip() == "":
				warnings.append(f"Thiếu hoặc chưa điền: {label}")

	elif doc_type == "report":
		# Support both old format (thong_tin_chung) and new scan_sxkd format (ho_ten at top level)
		info = data.get("thong_tin_chung", {})
		ho_ten = info.get("ho_ten") or data.get("ho_ten", "")
		vi_tri = info.get("vi_tri", "")
		phong_ban = info.get("phong_ban", "")

		if ho_ten: fields["Họ và tên"] = ho_ten.replace("Họ và tên:", "").strip()
		if vi_tri: fields["Vị trí công việc"] = vi_tri
		if phong_ban: fields["Phòng ban"] = phong_ban
		if data.get("tieu_de"): fields["Tiêu đề"] = data["tieu_de"]

	return {
		"fields": fields,
		"warnings": warnings,
	}


# ============================================================
# DIRECT FILE COMPLETENESS CHECKS
# ============================================================

# Placeholder patterns that count as "empty"
_PLACEHOLDER_CHARS = set(".-…_*/| \t")


# ============================================================
# VALIDATION FUNCTIONS
# ============================================================

# Data rows in the Excel template (STT 1-4)
XLSX_DATA_ROWS = [5, 8, 11, 14]
# ============================================================
# SHARED DAILY REPORT UTILITIES
# ============================================================


def parse_daily_report(file_url=None, raw_bytes=None, filename=None):
	"""Parse a daily report file and return text/JSON content.

	Handles .docx, .xlsx, .pdf. For .docx files that contain only
	embedded images (Teams/Zalo screenshots), falls back to
	GPT-4o Vision via ocr_daily_report_from_docx().

	Call with either:
	  - file_url: Frappe file URL (for evaluation.py / async jobs)
	  - raw_bytes + filename: raw content (for thu_viec.py / direct upload)

	Args:
		file_url: Frappe file URL string, or None.
		raw_bytes: bytes content of the file, or None.
		filename: original filename (used to detect extension), or None.

	Returns:
		String containing text or JSON from the daily report.
		Returns empty string on failure (never raises).
	"""
	import tempfile
	try:
		# ── Route by input type ──
		if file_url:
			# Frappe file URL path
			file_path = get_file_path_from_url(file_url)
			ext = os.path.splitext(file_path)[1].lower()
			if ext == ".pdf":
				content = parse_pdf(file_path)
				# Vision fallback for scanned/image-only PDFs
				if not content.strip():
					from cnb_2as.services.ocr_service import ocr_daily_report_from_pdf
					content = ocr_daily_report_from_pdf(file_path)
			elif ext in (".xlsx", ".xls"):
				content = parse_xlsx(file_path)
			else:
				content = parse_docx(file_path)
				# Vision fallback for image-only docx
				if not content.strip():
					from cnb_2as.services.ocr_service import ocr_daily_report_from_docx
					content = ocr_daily_report_from_docx(file_path)
			return content

		elif raw_bytes is not None and filename:
			# Raw bytes from multipart upload
			fname = filename.lower()
			if fname.endswith(".pdf"):
				# Write temp file for pdf parsing
				with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tf:
					tf.write(raw_bytes)
					tmp_path = tf.name
				try:
					content = parse_pdf(tmp_path)
					# Vision fallback for scanned/image-only PDFs
					if not content.strip():
						from cnb_2as.services.ocr_service import ocr_daily_report_from_pdf
						content = ocr_daily_report_from_pdf(tmp_path)
				finally:
					try:
						os.unlink(tmp_path)
					except Exception:
						pass
			elif fname.endswith((".xlsx", ".xls")):
				with tempfile.NamedTemporaryFile(suffix=os.path.splitext(fname)[1], delete=False) as tf:
					tf.write(raw_bytes)
					tmp_path = tf.name
				try:
					content = parse_xlsx(tmp_path)
				finally:
					try:
						os.unlink(tmp_path)
					except Exception:
						pass
			else:
				# .docx (default)
				with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tf:
					tf.write(raw_bytes)
					tmp_path = tf.name
				try:
					content = parse_docx(tmp_path)
					# Vision fallback for image-only docx
					if not content.strip():
						from cnb_2as.services.ocr_service import ocr_daily_report_from_docx
						content = ocr_daily_report_from_docx(tmp_path)
				finally:
					try:
						os.unlink(tmp_path)
					except Exception:
						pass
			return content

		return ""
	except Exception as e:
		print(f"[parse_daily_report] Error: {str(e)[:200]}")
		return ""


# ── Re-exports for backwards compatibility ─────────────────────────────────────
# These functions were moved to cnb_2as.services.validators but are re-exported
# here so existing imports (e.g. in evaluation.py) continue to work without changes.
from cnb_2as.services.validators import (  # noqa: E402
    check_eval_docx_completeness,
    check_report_xlsx_completeness,
    count_working_days,
    validate_cross_document_consistency,
    validate_xlsx_kpi_ratio,
    validate_xlsx_product_links,
)
