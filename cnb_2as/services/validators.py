# Copyright (c) 2025, antruong and contributors
# For license information, please see license.txt

"""Document and file validators for HR evaluation workflows.

Validates structure and content of DOCX/XLSX evaluation files,
checks completeness, product links, KPI ratios, and cross-document
consistency between Word and Excel files.

Imported by: api/evaluation.py, api/thu_viec.py
"""

import re
from datetime import datetime, date, timedelta

import frappe

# Needed for reading documents in validators
from docx import Document as DocxDocument
import openpyxl

from cnb_2as.services.document_parser import get_file_path_from_url, parse_docx, parse_xlsx


def _check_field(text, fields, warnings, field_name, keywords, checklist=None):
	"""Check if a field exists in the text and extract its value.

	Args:
		text: Full text content.
		fields: Dict to store extracted field values.
		warnings: List to append warnings to.
		field_name: Human-readable field name.
		keywords: List of keyword variations to search for.
		checklist: Optional list to append {"label": string, "is_filled": bool}
	"""
	found = False
	for kw in keywords:
		idx = text.find(kw)
		if idx >= 0:
			# Try to extract value after the keyword
			after = text[idx + len(kw):idx + len(kw) + 200]
			# Look for value after | or : separator
			value = ""
			for sep in ["|", ":"]:
				sep_idx = after.find(sep)
				if sep_idx >= 0:
					val_part = after[sep_idx + 1:]
					# Get until next | or newline
					end_idx = len(val_part)
					for end_char in ["|", "\n"]:
						ei = val_part.find(end_char)
						if ei >= 0 and ei < end_idx:
							end_idx = ei
					value = val_part[:end_idx].strip()
					break

			if value and value not in ["", "...", "......", "Giá trị"]:
				fields[field_name] = value
				found = True
			break

	if not found:
		warnings.append(f"Thiếu hoặc chưa điền: {field_name}")
	
	if checklist is not None:
		checklist.append({"label": field_name, "is_filled": found})


def _check_signature(text, warnings, signer_role, checklist=None):
	"""Check if a signature section has been signed.

	Args:
		text: Full text content.
		warnings: List to append warnings to.
		signer_role: Role name to check for signature.
		checklist: Optional list to append status
	"""
	lower_text = text.lower()
	role_lower = signer_role.lower()

	is_filled = True
	if role_lower in lower_text:
		# Check if there's a "đã ký" or actual name near the role
		role_idx = lower_text.find(role_lower)
		context = lower_text[role_idx:role_idx + 300]
		if "chưa ký" in context or "chưa có" in context:
			warnings.append(f"Chưa ký: {signer_role}")
			is_filled = False
		elif "đã ký" not in context and "ký" not in context:
			warnings.append(f"Chưa xác nhận chữ ký: {signer_role}")
			is_filled = False
	else:
		warnings.append(f"Thiếu phần chữ ký: {signer_role}")
		is_filled = False

	if checklist is not None:
		checklist.append({"label": f"Chữ ký: {signer_role}", "is_filled": is_filled})


def _check_content_section(text, warnings, section_keyword, warning_label, checklist=None):
	"""Check if a content section has actual content filled in.

	Args:
		text: Full text content.
		warnings: List to append warnings to.
		section_keyword: Keyword to find the section.
		warning_label: Human-readable label for warnings.
		checklist: Optional list to append status
	"""
	idx = text.lower().find(section_keyword.lower())
	if idx < 0:
		warnings.append(f"Thiếu phần: {warning_label}")
		if checklist is not None:
			checklist.append({"label": warning_label, "is_filled": False})
		return  # Section not found, might not be applicable

	after = text[idx + len(section_keyword):idx + len(section_keyword) + 500]
	# Check if content is just dots or empty
	content_lines = [
		line.strip()
		for line in after.split("\n")[:5]
		if line.strip()
		and not line.strip().startswith("#")
		and not line.strip().startswith("---")
	]

	has_real_content = False
	for line in content_lines:
		# Filter out placeholder patterns
		cleaned = line.replace(".", "").replace("…", "").replace("-", "").replace("*", "").strip()
		if len(cleaned) > 5:
			has_real_content = True
			break

	if not has_real_content:
		warnings.append(f"Chưa điền nội dung: {warning_label}")
	
	if checklist is not None:
		checklist.append({"label": warning_label, "is_filled": has_real_content})


def _is_cell_empty(value):
	"""Check if a cell value is effectively empty (None, blank, or placeholder)."""
	if value is None:
		return True
	s = str(value).strip()
	if not s:
		return True
	# Check for placeholder-only content like "...", "___", "…...", "----"
	cleaned = "".join(c for c in s if c not in _PLACEHOLDER_CHARS)
	if len(cleaned) <= 2:
		return True
	return False


def check_eval_docx_completeness(file_url):
	"""Check Word evaluation form for empty fields by reading table structure directly.

	Reads the .docx file's table cells to detect truly empty fields.

	Args:
		file_url: Frappe file URL for the Word evaluation form.

	Returns:
		List of warning strings for empty/missing fields.
	"""
	file_path = get_file_path_from_url(file_url)
	ext = os.path.splitext(file_path)[1].lower()
	if ext != ".docx":
		return []

	try:
		doc = DocxDocument(file_path)
	except Exception:
		return ["Không thể đọc file Word để kiểm tra"]

	warnings = []
	tables = doc.tables

	if not tables:
		warnings.append("File Đánh giá tái ký: Không tìm thấy bảng biểu nào")
		return warnings

	# --- Table 0: Thông tin nhân viên ---
	if len(tables) > 0:
		info_table = tables[0]
		# Scan all cells in the info table for key fields
		info_text = ""
		for row in info_table.rows:
			for cell in row.cells:
				info_text += " " + cell.text.strip()

		info_lower = info_text.lower()
		info_fields = {
			"Họ và tên": ["họ và tên", "họ tên"],
			"MSNV": ["msnv", "mã số nhân viên"],
			"Vị trí công việc": ["vị trí công việc", "chức danh"],
			"Phòng ban": ["phòng ban", "phòng – ban"],
			"Ngày bắt đầu HĐ": ["ngày bắt đầu hđ", "ngày nhận việc"],
			"Ngày hết hạn HĐ": ["ngày hết hạn hđ", "ngày hết hạn"],
		}

		for label, keywords in info_fields.items():
			found_val = False
			for kw in keywords:
				idx = info_lower.find(kw)
				if idx >= 0:
					# Check if there's real content after the keyword
					after = info_text[idx + len(kw):idx + len(kw) + 100]
					# Remove separators and check
					for sep in [":", "|"]:
						si = after.find(sep)
						if si >= 0:
							val = after[si + 1:].split("|")[0].split("\n")[0].strip()
							if not _is_cell_empty(val):
								found_val = True
								break
					break
			if not found_val:
				warnings.append(f"File Đánh giá tái ký: Chưa điền '{label}'")

	# --- Scan all tables for key sections ---
	all_cell_text = []
	for table in tables:
		for row in table.rows:
			row_cells = []
			for cell in row.cells:
				row_cells.append(cell.text.strip())
			all_cell_text.append(row_cells)

	# Helper: find rows containing a keyword and check if adjacent cells have content
	def _find_section_filled(keyword, check_cols=None):
		"""Search all table cells for keyword and check if associated value cells are filled."""
		kw_lower = keyword.lower()
		for row_cells in all_cell_text:
			for ci, cell_val in enumerate(row_cells):
				if kw_lower in cell_val.lower():
					# Check subsequent cells in same row
					for check_ci in range(ci + 1, len(row_cells)):
						if not _is_cell_empty(row_cells[check_ci]):
							return True
					# Also check if same cell contains value after keyword
					after_kw = cell_val[cell_val.lower().find(kw_lower) + len(keyword):]
					if not _is_cell_empty(after_kw):
						return True
		return False

	# Mục III. Nhận xét
	section_checks = [
		("Nhận xét chi tiết kết quả công việc", "Nhận xét chi tiết kết quả công việc"),
		("Ưu điểm của CBNV", "Ưu điểm của CBNV"),
		("Hạn chế của CBNV", "Hạn chế của CBNV"),
		("Giải pháp, yêu cầu khắc phục", "Giải pháp/yêu cầu khắc phục"),
		("Đề xuất của Quản lý trực tiếp", "Đề xuất của Quản lý trực tiếp"),
		("Đề xuất của", "Đề xuất của Lãnh đạo Ban"),
	]

	for keyword, label in section_checks:
		if not _find_section_filled(keyword):
			warnings.append(f"File Đánh giá tái ký: Chưa điền '{label}'")

	return warnings


def check_report_xlsx_completeness(file_url):
	"""Check Excel work report for empty fields by reading cell values directly.

	Reads the .xlsx file's cells to detect truly empty data.

	Args:
		file_url: Frappe file URL for the Excel work report.

	Returns:
		List of warning strings for empty/missing fields.
	"""
	file_path = get_file_path_from_url(file_url)
	ext = os.path.splitext(file_path)[1].lower()
	if ext != ".xlsx":
		return []

	try:
		wb = load_workbook(file_path, read_only=True, data_only=True)
	except Exception:
		return ["Không thể đọc file Excel để kiểm tra"]

	warnings = []

	# Use the last sheet (most recent month) or active sheet
	ws = wb.active
	if len(wb.sheetnames) > 1:
		ws = wb[wb.sheetnames[-1]]

	# Read all rows to find structure
	all_rows = []
	for row in ws.iter_rows(values_only=True):
		all_rows.append([str(c).strip() if c is not None else "" for c in row])

	wb.close()

	if not all_rows:
		warnings.append("File Báo cáo kết quả: File rỗng")
		return warnings

	# Helper: find row index containing keyword
	def _find_row_with(keyword):
		kw = keyword.lower()
		for ri, row in enumerate(all_rows):
			full = " ".join(row).lower()
			if kw in full:
				return ri
		return -1

	# 1. Check "Họ và tên"
	name_row = _find_row_with("họ và tên")
	if name_row >= 0:
		row_text = " ".join(all_rows[name_row])
		# Remove the label and check if a name exists
		cleaned = row_text.lower().replace("họ và tên:", "").replace("họ và tên", "").strip()
		if _is_cell_empty(cleaned):
			warnings.append("File Báo cáo kết quả: Chưa điền 'Họ và tên'")

	# 2. Check work rows - find header row with "TỶ TRỌNG"
	header_row = _find_row_with("tỷ trọng")
	if header_row >= 0:
		# Data rows follow the header. Scan for numeric STT rows
		for ri in range(header_row + 1, min(header_row + 20, len(all_rows))):
			row = all_rows[ri]
			if not row or not row[0]:
				continue
			# Check if first cell is a number (STT)
			try:
				stt = int(float(row[0]))
			except (ValueError, TypeError):
				# Check for "TỶ LỆ ĐẠT" row = end of section
				if "tỷ lệ đạt" in " ".join(row).lower():
					break
				continue

			task_name = row[1] if len(row) > 1 else f"Mục {stt}"
			task_short = str(task_name).strip()[:40]

			# Check Tỷ trọng (column index varies, usually col 3-4)
			# Check KPI result (column index varies, usually col 7-8)
			# We look for empty numeric cells in the row
			has_weight = False
			has_kpi_result = False
			for ci in range(2, min(len(row), 10)):
				val = row[ci]
				try:
					f = float(val)
					if 0 < f <= 1:
						has_weight = True
					if f > 0:
						has_kpi_result = True
				except (ValueError, TypeError):
					pass

			if not has_weight:
				warnings.append(f"File Báo cáo kết quả: Mục '{task_short}' — Chưa điền Tỷ trọng/KPI")

	# 3. Check BOD xét duyệt
	bod_row = _find_row_with("bod")
	hod_row = _find_row_with("hod")

	# Check in PHẦN TRÌNH VÀ XÉT DUYỆT section
	review_row = _find_row_with("phần trình và xét duyệt")
	if review_row >= 0:
		# HOD and BOD rows should be after this
		for ri in range(review_row + 1, min(review_row + 6, len(all_rows))):
			row = all_rows[ri]
			row_text = " ".join(row).lower()
			if "hod" in row_text:
				# Check if there's content besides "HOD"
				content = row_text.replace("hod", "").replace("xét duyệt", "").replace("ý kiến", "").strip()
				if _is_cell_empty(content):
					warnings.append("File Báo cáo kết quả: Chưa điền 'Ý kiến HOD'")
			elif "bod" in row_text:
				content = row_text.replace("bod", "").replace("xét duyệt", "").replace("ý kiến", "").strip()
				if _is_cell_empty(content):
					warnings.append("File Báo cáo kết quả: Chưa điền 'Ý kiến BOD'")

	# 4. Check Ranking
	ranking_row = _find_row_with("ranking")
	if ranking_row >= 0:
		row = all_rows[ranking_row]
		# Find the Ranking column value
		for ci, val in enumerate(row):
			if "ranking" in val.lower():
				# Check adjacent cells for a value
				has_ranking = False
				for check_ci in range(ci + 1, len(row)):
					if not _is_cell_empty(row[check_ci]):
						has_ranking = True
						break
				if not has_ranking:
					warnings.append("File Báo cáo kết quả: Chưa điền 'Ranking'")
				break

	return warnings


def validate_xlsx_product_links(file_url):
	"""Validate that product link cells (column J) are not empty.

	Checks column 10 (J) in the Excel work report for each data row.

	Args:
		file_url: Frappe file URL for the Excel file.

	Returns:
		List of warning dicts, empty if all cells are filled.
	"""
	file_path = get_file_path_from_url(file_url)
	ext = os.path.splitext(file_path)[1].lower()
	if ext != ".xlsx":
		return []

	wb = load_workbook(file_path, read_only=True, data_only=True)
	ws = wb.active
	warnings = []

	for row_num in XLSX_DATA_ROWS:
		cell_val = ws.cell(row=row_num, column=10).value
		task_name = ws.cell(row=row_num, column=2).value or f"STT {row_num}"
		if cell_val is None or str(cell_val).strip() == "":
			warnings.append({
				"type": "missing_product_link",
				"row": row_num,
				"task": str(task_name).strip(),
				"message": (
					f"Dòng {row_num} ({str(task_name).strip()[:40]}): "
					f"Cột 'Link sản phẩm' đang để trống"
				),
			})

	wb.close()
	return warnings


def validate_xlsx_kpi_ratio(file_url):
	"""Validate that KPI ratio cells (column G) are not empty.

	Checks column 7 (G = Tỷ lệ KPI in Kết quả thực hiện) for each data row.

	Args:
		file_url: Frappe file URL for the Excel file.

	Returns:
		List of warning dicts, empty if all cells are filled.
	"""
	file_path = get_file_path_from_url(file_url)
	ext = os.path.splitext(file_path)[1].lower()
	if ext != ".xlsx":
		return []

	wb = load_workbook(file_path, read_only=True, data_only=True)
	ws = wb.active
	warnings = []

	for row_num in XLSX_DATA_ROWS:
		cell_val = ws.cell(row=row_num, column=7).value
		task_name = ws.cell(row=row_num, column=2).value or f"STT {row_num}"
		if cell_val is None or str(cell_val).strip() == "":
			warnings.append({
				"type": "missing_kpi_ratio",
				"row": row_num,
				"task": str(task_name).strip(),
				"message": (
					f"Dòng {row_num} ({str(task_name).strip()[:40]}): "
					f"Cột 'Tỷ lệ KPI' (Kết quả thực hiện) đang để trống"
				),
			})

	wb.close()
	return warnings


def _extract_docx_tasks(file_path):
	"""Extract task entries from DOCX Table 2 for consistency checking.

	Returns list of dicts with task name and self-assessment info.
	"""
	doc = DocxDocument(file_path)
	tasks = []

	if len(doc.tables) < 3:
		return tasks

	table = doc.tables[2]
	# Skip header rows (0, 1), data starts at row 2
	for row_idx in range(2, len(table.rows)):
		row = table.rows[row_idx]
		cells = [cell.text.strip() for cell in row.cells]
		if not cells or not cells[0]:
			continue
		task_info = {
			"name": cells[0][:100],
			"level": cells[1] if len(cells) > 1 else "",
			"rate": cells[2] if len(cells) > 2 else "",
		}
		tasks.append(task_info)

	return tasks


def _extract_xlsx_tasks(file_path):
	"""Extract task entries from Excel for consistency checking.

	Returns list of dicts with task name and KPI info.
	"""
	wb = load_workbook(file_path, read_only=True, data_only=True)
	ws = wb.active
	tasks = []

	for row_num in XLSX_DATA_ROWS:
		task_name = ws.cell(row=row_num, column=2).value
		if task_name and str(task_name).strip():
			kpi_ratio = ws.cell(row=row_num, column=7).value
			weight = ws.cell(row=row_num, column=4).value
			tasks.append({
				"name": str(task_name).strip()[:100],
				"row": row_num,
				"kpi_ratio": kpi_ratio,
				"weight": weight,
			})

	wb.close()
	return tasks


def _fuzzy_name_match(name_a, name_b):
	"""Check if two task names are similar enough to be considered matching.

	Uses simple substring matching — checks if the shorter name's first
	20 characters appear within the longer name.
	"""
	a = name_a.lower().strip()
	b = name_b.lower().strip()
	if a == b:
		return True
	short = a if len(a) <= len(b) else b
	long_str = b if len(a) <= len(b) else a
	# Check first 20 chars of shorter name in longer
	key = short[:20]
	return key in long_str


def validate_cross_document_consistency(eval_file_url, report_file_url):
	"""Compare task data between Word and Excel to detect inconsistencies.

	Args:
		eval_file_url: Frappe file URL for the Word evaluation form.
		report_file_url: Frappe file URL for the Excel work report.

	Returns:
		Dict with is_consistent flag, warnings list, and errors list.
	"""
	result = {"is_consistent": True, "warnings": [], "errors": []}

	eval_path = get_file_path_from_url(eval_file_url)
	report_path = get_file_path_from_url(report_file_url)

	eval_ext = os.path.splitext(eval_path)[1].lower()
	report_ext = os.path.splitext(report_path)[1].lower()

	if eval_ext != ".docx" or report_ext != ".xlsx":
		result["warnings"].append({
			"type": "format_skip",
			"message": (
				"Kiểm tra đồng nhất chỉ hỗ trợ Word (.docx) + Excel (.xlsx). "
				"Bỏ qua kiểm tra."
			),
		})
		return result

	docx_tasks = _extract_docx_tasks(eval_path)
	xlsx_tasks = _extract_xlsx_tasks(report_path)

	# Check 1: Number of tasks
	if len(docx_tasks) != len(xlsx_tasks):
		result["warnings"].append({
			"type": "task_count_mismatch",
			"message": (
				f"Số lượng mục công việc không khớp: "
				f"Word có {len(docx_tasks)} mục, Excel có {len(xlsx_tasks)} mục"
			),
		})
		result["is_consistent"] = False

	# Check 2: Task name matching
	min_count = min(len(docx_tasks), len(xlsx_tasks))
	for i in range(min_count):
		docx_name = docx_tasks[i]["name"]
		xlsx_name = xlsx_tasks[i]["name"]
		if not _fuzzy_name_match(docx_name, xlsx_name):
			result["warnings"].append({
				"type": "task_name_mismatch",
				"message": (
					f"Mục {i + 1}: Tên công việc không khớp giữa Word và Excel.\n"
					f"  Word: \"{docx_name[:50]}...\"\n"
					f"  Excel: \"{xlsx_name[:50]}...\""
				),
			})
			result["is_consistent"] = False

	# Check 3: Self-assessment rate vs KPI ratio
	for i in range(min_count):
		docx_rate_str = docx_tasks[i].get("rate", "").replace("%", "").strip()
		xlsx_kpi = xlsx_tasks[i].get("kpi_ratio")

		if not docx_rate_str or xlsx_kpi is None:
			continue

		try:
			docx_rate = float(docx_rate_str) / 100.0
			xlsx_kpi_float = float(xlsx_kpi)
			# If difference > 30%, flag as suspicious
			if abs(docx_rate - xlsx_kpi_float) > 0.30:
				result["warnings"].append({
					"type": "rate_mismatch",
					"message": (
						f"Mục {i + 1} ({xlsx_tasks[i]['name'][:30]}...): "
						f"Tỷ lệ đánh giá chênh lệch lớn — "
						f"Word: {docx_rate_str}%, Excel KPI: "
						f"{xlsx_kpi_float * 100:.0f}%"
					),
				})
				result["is_consistent"] = False
		except (ValueError, TypeError):
			pass

	return result


def count_working_days(ngay_bd, ngay_kt):
	"""Count Monday-Friday working days between two dates (inclusive).

	Args:
		ngay_bd: Start date string in ISO format (YYYY-MM-DD).
		ngay_kt: End date string in ISO format (YYYY-MM-DD).

	Returns:
		Int number of working days (0 if dates invalid or empty).
	"""
	if not ngay_bd or not ngay_kt:
		return 0
	from datetime import date, timedelta
	try:
		d_start = date.fromisoformat(ngay_bd)
		d_end = date.fromisoformat(ngay_kt)
	except ValueError:
		return 0
	if d_end < d_start:
		return 0
	count = 0
	cur = d_start
	while cur <= d_end:
		if cur.weekday() < 5:  # 0=Mon … 4=Fri
			count += 1
		cur += timedelta(days=1)
	return count
