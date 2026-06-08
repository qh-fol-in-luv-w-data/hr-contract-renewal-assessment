"""
OCR Báo cáo kết quả công việc: Scan PDF → Excel
=================================================

Hàm này dùng OpenAI GPT-4o Vision để đọc file PDF scan (Báo cáo kết quả công việc)
và xuất ra file Excel có cấu trúc.

Yêu cầu cài đặt:
    pip install openai PyMuPDF Pillow openpyxl

Cách dùng:
    from ocr_report_to_excel import convert_report_pdf_to_excel
    
    output_path = convert_report_pdf_to_excel(
        pdf_path="path/to/bao_cao.pdf",
        api_key="sk-xxx",
        output_path="output.xlsx"  # optional
    )

Hoặc chạy trực tiếp:
    python ocr_report_to_excel.py path/to/bao_cao.pdf --api-key sk-xxx
"""

import base64
import io
import json
import os
import sys
import time


def pdf_to_images(pdf_path, dpi=200, max_pages=20):
    """Chuyển PDF thành danh sách ảnh base64 (JPEG).
    
    Args:
        pdf_path: Đường dẫn file PDF.
        dpi: Độ phân giải render (mặc định 200).
        max_pages: Số trang tối đa xử lý.
    
    Returns:
        List[str]: Danh sách chuỗi base64 JPEG.
    """
    import fitz  # PyMuPDF
    from PIL import Image
    
    doc = fitz.open(pdf_path)
    images = []
    
    zoom = dpi / 72.0
    mat = fitz.Matrix(zoom, zoom)
    
    for page_num in range(min(len(doc), max_pages)):
        page = doc[page_num]
        pix = page.get_pixmap(matrix=mat)
        
        img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
        
        # Resize nếu quá lớn (max 2048px)
        max_dim = 2048
        if max(img.size) > max_dim:
            ratio = max_dim / max(img.size)
            new_size = (int(img.size[0] * ratio), int(img.size[1] * ratio))
            img = img.resize(new_size)
        
        buffer = io.BytesIO()
        img.save(buffer, format="JPEG", quality=85)
        b64 = base64.b64encode(buffer.getvalue()).decode("utf-8")
        images.append(b64)
    
    doc.close()
    print(f"[OCR] Đã convert {len(images)} trang từ PDF")
    return images


def ocr_pdf_with_vision(pdf_path, api_key, model="gpt-4o"):
    """Đọc OCR file PDF bằng OpenAI GPT-4o Vision, trả về JSON có cấu trúc.
    
    Args:
        pdf_path: Đường dẫn file PDF scan.
        api_key: OpenAI API key.
        model: Model OpenAI (mặc định gpt-4o).
    
    Returns:
        dict: Dữ liệu đã trích xuất dưới dạng JSON.
    """
    from openai import OpenAI
    
    client = OpenAI(api_key=api_key)
    
    # Convert PDF → ảnh
    base64_images = pdf_to_images(pdf_path)
    
    if not base64_images:
        raise ValueError("Không thể chuyển PDF thành hình ảnh")
    
    # Tạo prompt yêu cầu GPT-4o trích xuất dữ liệu dạng JSON
    system_prompt = """Bạn là chuyên gia OCR và trích xuất dữ liệu từ tài liệu scan.
Nhiệm vụ: Đọc file "Báo cáo kết quả công việc" và trích xuất TOÀN BỘ nội dung thành JSON.

Trả về JSON có cấu trúc như sau:
{
    "thong_tin_chung": {
        "ho_ten": "...",
        "msnv": "...", 
        "phong_ban": "...",
        "vi_tri": "...",
        "thang_danh_gia": "...",
        "ky_danh_gia": "..."
    },
    "bang_cong_viec": [
        {
            "stt": 1,
            "noi_dung_cong_viec": "...",
            "ket_qua": "...",
            "ty_trong": "...",
            "tu_danh_gia": "...",
            "quan_ly_danh_gia": "...",
            "ghi_chu": "..."
        }
    ],
    "tong_ket": {
        "tong_ty_trong": "...",
        "diem_tu_danh_gia": "...",
        "diem_quan_ly": "...",
        "nhan_xet_chung": "...",
        "de_xuat": "..."
    },
    "noi_dung_khac": "... (bất kỳ nội dung nào khác trong tài liệu)"
}

QUY TẮC:
- Đọc CHÍNH XÁC nội dung trong tài liệu, KHÔNG thêm bớt
- Nếu một trường không có dữ liệu, ghi ""
- Nếu bảng có nhiều cột khác, thêm vào object tương ứng
- Giữ nguyên tiếng Việt, KHÔNG dịch
- Nếu cấu trúc tài liệu khác mẫu trên, hãy TỰ ĐIỀU CHỈNH cấu trúc JSON cho phù hợp
- Trả về JSON thuần, KHÔNG wrap trong code block"""

    # Build content với ảnh
    content = [{"type": "text", "text": "Hãy trích xuất toàn bộ nội dung từ tài liệu scan sau:"}]
    
    for b64_img in base64_images:
        content.append({
            "type": "image_url",
            "image_url": {
                "url": f"data:image/jpeg;base64,{b64_img}",
                "detail": "high",
            },
        })
    
    print(f"[OCR] Đang gọi OpenAI {model} với {len(base64_images)} ảnh...")
    start = time.time()
    
    response = client.chat.completions.create(
        model=model,
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": content},
        ],
        response_format={"type": "json_object"},
        max_tokens=16384,
        temperature=0.1,
    )
    
    raw = response.choices[0].message.content
    elapsed = time.time() - start
    tokens = response.usage.total_tokens if response.usage else 0
    print(f"[OCR] Hoàn tất: {elapsed:.1f}s, {tokens} tokens")
    
    # Parse JSON
    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        # Nếu GPT wrap trong code block
        clean = raw.strip()
        if clean.startswith("```"):
            clean = clean.split("\n", 1)[1] if "\n" in clean else clean[3:]
        if clean.endswith("```"):
            clean = clean[:-3]
        data = json.loads(clean.strip())
    
    return data


def json_to_excel(data, output_path):
    """Chuyển dữ liệu JSON thành file Excel.
    
    Args:
        data: Dict dữ liệu đã trích xuất từ OCR.
        output_path: Đường dẫn file Excel đầu ra.
    
    Returns:
        str: Đường dẫn file Excel đã tạo.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Báo cáo kết quả công việc"
    
    # Styles
    header_font = Font(name="Arial", bold=True, size=12, color="FFFFFF")
    header_fill = PatternFill(start_color="7C3AED", end_color="7C3AED", fill_type="solid")
    sub_header_font = Font(name="Arial", bold=True, size=10)
    sub_header_fill = PatternFill(start_color="E8DEF8", end_color="E8DEF8", fill_type="solid")
    normal_font = Font(name="Arial", size=10)
    bold_font = Font(name="Arial", bold=True, size=10)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    wrap_alignment = Alignment(wrap_text=True, vertical="top")
    center_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    row = 1
    
    # === THÔNG TIN CHUNG ===
    thong_tin = data.get("thong_tin_chung", {})
    if thong_tin:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        cell = ws.cell(row=row, column=1, value="THÔNG TIN CHUNG")
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_alignment
        row += 1
        
        for key, value in thong_tin.items():
            label = key.replace("_", " ").title()
            ws.cell(row=row, column=1, value=label).font = bold_font
            ws.cell(row=row, column=1).border = thin_border
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
            ws.cell(row=row, column=2, value=str(value)).font = normal_font
            ws.cell(row=row, column=2).border = thin_border
            row += 1
        
        row += 1
    
    # === BẢNG CÔNG VIỆC ===
    bang_cv = data.get("bang_cong_viec", [])
    if bang_cv:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        cell = ws.cell(row=row, column=1, value="BẢNG CÔNG VIỆC")
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_alignment
        row += 1
        
        # Detect headers from first item
        if bang_cv:
            headers = list(bang_cv[0].keys())
            # Header row
            for col_idx, h in enumerate(headers, 1):
                cell = ws.cell(row=row, column=col_idx, value=h.replace("_", " ").upper())
                cell.font = sub_header_font
                cell.fill = sub_header_fill
                cell.border = thin_border
                cell.alignment = center_alignment
            row += 1
            
            # Data rows
            for item in bang_cv:
                for col_idx, h in enumerate(headers, 1):
                    cell = ws.cell(row=row, column=col_idx, value=str(item.get(h, "")))
                    cell.font = normal_font
                    cell.border = thin_border
                    cell.alignment = wrap_alignment
                row += 1
        
        row += 1
    
    # === TỔNG KẾT ===
    tong_ket = data.get("tong_ket", {})
    if tong_ket:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        cell = ws.cell(row=row, column=1, value="TỔNG KẾT")
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_alignment
        row += 1
        
        for key, value in tong_ket.items():
            label = key.replace("_", " ").title()
            ws.cell(row=row, column=1, value=label).font = bold_font
            ws.cell(row=row, column=1).border = thin_border
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
            ws.cell(row=row, column=2, value=str(value)).font = normal_font
            ws.cell(row=row, column=2).border = thin_border
            row += 1
        
        row += 1
    
    # === NỘI DUNG KHÁC ===
    noi_dung_khac = data.get("noi_dung_khac", "")
    if noi_dung_khac:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        cell = ws.cell(row=row, column=1, value="NỘI DUNG KHÁC")
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_alignment
        row += 1
        
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        ws.cell(row=row, column=1, value=str(noi_dung_khac)).font = normal_font
        ws.cell(row=row, column=1).alignment = wrap_alignment
        row += 1
    
    # === Handle extra keys not in template ===
    known_keys = {"thong_tin_chung", "bang_cong_viec", "tong_ket", "noi_dung_khac"}
    for key in data:
        if key in known_keys:
            continue
        
        value = data[key]
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        cell = ws.cell(row=row, column=1, value=key.replace("_", " ").upper())
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_alignment
        row += 1
        
        if isinstance(value, list):
            # Nếu là danh sách, tạo bảng
            if value and isinstance(value[0], dict):
                headers = list(value[0].keys())
                for col_idx, h in enumerate(headers, 1):
                    cell = ws.cell(row=row, column=col_idx, value=h.replace("_", " ").upper())
                    cell.font = sub_header_font
                    cell.fill = sub_header_fill
                    cell.border = thin_border
                row += 1
                for item in value:
                    for col_idx, h in enumerate(headers, 1):
                        cell = ws.cell(row=row, column=col_idx, value=str(item.get(h, "")))
                        cell.font = normal_font
                        cell.border = thin_border
                        cell.alignment = wrap_alignment
                    row += 1
        elif isinstance(value, dict):
            for k, v in value.items():
                ws.cell(row=row, column=1, value=k.replace("_", " ").title()).font = bold_font
                ws.cell(row=row, column=1).border = thin_border
                ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
                ws.cell(row=row, column=2, value=str(v)).font = normal_font
                ws.cell(row=row, column=2).border = thin_border
                row += 1
        else:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
            ws.cell(row=row, column=1, value=str(value)).font = normal_font
            ws.cell(row=row, column=1).alignment = wrap_alignment
            row += 1
        
        row += 1
    
    # Auto-fit column widths
    for col in ws.columns:
        max_length = 0
        column_letter = col[0].column_letter
        for cell in col:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except Exception:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = max(adjusted_width, 12)
    
    wb.save(output_path)
    print(f"[Excel] Đã lưu: {output_path}")
    return output_path


def convert_report_pdf_to_excel(pdf_path, api_key, output_path=None, model="gpt-4o"):
    """Hàm chính: Chuyển PDF scan Báo cáo công việc → Excel.
    
    Args:
        pdf_path: Đường dẫn file PDF scan đầu vào.
        api_key: OpenAI API key (sk-xxx).
        output_path: Đường dẫn file Excel đầu ra (optional).
                     Nếu không truyền, tự tạo tên từ file PDF.
        model: OpenAI model (mặc định gpt-4o).
    
    Returns:
        str: Đường dẫn file Excel đã tạo.
    
    Raises:
        ValueError: Nếu file không hợp lệ hoặc OCR thất bại.
        FileNotFoundError: Nếu file PDF không tồn tại.
    
    Ví dụ:
        >>> path = convert_report_pdf_to_excel(
        ...     "bao_cao_cong_viec.pdf",
        ...     api_key="sk-proj-xxx"
        ... )
        >>> print(f"File Excel: {path}")
    """
    # Validate input
    if not os.path.exists(pdf_path):
        raise FileNotFoundError(f"File không tồn tại: {pdf_path}")
    
    ext = os.path.splitext(pdf_path)[1].lower()
    if ext != ".pdf":
        raise ValueError(f"Chỉ hỗ trợ file PDF, file hiện tại: {ext}")
    
    # Generate output path if not provided
    if not output_path:
        base = os.path.splitext(pdf_path)[0]
        output_path = f"{base}_output.xlsx"
    
    print(f"\n{'='*60}")
    print(f"  OCR Báo cáo công việc: PDF → Excel")
    print(f"  Input:  {pdf_path}")
    print(f"  Output: {output_path}")
    print(f"  Model:  {model}")
    print(f"{'='*60}\n")
    
    start = time.time()
    
    # Step 1: OCR bằng OpenAI Vision
    data = ocr_pdf_with_vision(pdf_path, api_key, model)
    
    # Step 2: Chuyển JSON → Excel
    result = json_to_excel(data, output_path)
    
    total = time.time() - start
    print(f"\n✅ Hoàn tất trong {total:.1f}s")
    print(f"📄 File Excel: {result}")
    
    return result


# ============================================================
# CLI: chạy trực tiếp từ terminal
# ============================================================
if __name__ == "__main__":
    import argparse
    
    parser = argparse.ArgumentParser(
        description="OCR Báo cáo kết quả công việc: PDF scan → Excel"
    )
    parser.add_argument("pdf_path", help="Đường dẫn file PDF scan đầu vào")
    parser.add_argument("--api-key", help="OpenAI API key", 
                        default=os.environ.get("OPENAI_API_KEY"))
    parser.add_argument("--output", "-o", help="Đường dẫn file Excel đầu ra")
    parser.add_argument("--model", default="gpt-4o", help="OpenAI model (default: gpt-4o)")
    
    args = parser.parse_args()
    
    if not args.api_key:
        print("❌ Thiếu API key. Dùng --api-key hoặc set OPENAI_API_KEY env var")
        sys.exit(1)
    
    try:
        convert_report_pdf_to_excel(
            pdf_path=args.pdf_path,
            api_key=args.api_key,
            output_path=args.output,
            model=args.model,
        )
    except Exception as e:
        print(f"\n❌ Lỗi: {e}")
        sys.exit(1)
